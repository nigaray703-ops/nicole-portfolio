from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "artifacts" / "Nicole_Nikareayi_Public_CV.docx"

NAVY = RGBColor(18, 56, 74)
TEAL = RGBColor(0, 126, 128)
MUTED = RGBColor(76, 101, 114)
INK = RGBColor(31, 43, 49)


def set_run_font(run, size, color=INK, bold=False, italic=False):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic


def set_cell_free_paragraph(paragraph, before=0, after=0, line=1.0):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def add_hyperlink(paragraph, text, url, size=9.2):
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "007E80")
    properties.append(color)
    size_element = OxmlElement("w:sz")
    size_element.set(qn("w:val"), str(int(size * 2)))
    properties.append(size_element)
    size_cs = OxmlElement("w:szCs")
    size_cs.set(qn("w:val"), str(int(size * 2)))
    properties.append(size_cs)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    properties.append(fonts)
    run.append(properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_section_heading(document, text):
    paragraph = document.add_paragraph(style="Heading 1")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text.upper())
    set_run_font(run, 10.5, TEAL, bold=True)
    return paragraph


def add_role_header(document, title, organisation, dates):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.line_spacing = 1.0

    title_run = paragraph.add_run(title)
    set_run_font(title_run, 10, NAVY, bold=True)

    organisation_run = paragraph.add_run(f" | {organisation}")
    set_run_font(organisation_run, 9.4, INK)

    date_run = paragraph.add_run(f" | {dates}")
    set_run_font(date_run, 9.2, MUTED, italic=True)


def add_body(document, text, after=3, bold=False):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(text)
    set_run_font(run, 9.2, INK, bold=bold)
    return paragraph


def add_bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.2)
    paragraph.paragraph_format.first_line_indent = Inches(-0.13)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(1.5)
    paragraph.paragraph_format.line_spacing = 1.04
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(text)
    set_run_font(run, 8.9, INK)
    return paragraph


def add_footer(section):
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("Nicole Nikareayi | Public CV | ")
    set_run_font(run, 8, MUTED)

    page_field = OxmlElement("w:fldSimple")
    page_field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(page_field)

    separator = paragraph.add_run(" / ")
    set_run_font(separator, 8, MUTED)

    pages_field = OxmlElement("w:fldSimple")
    pages_field.set(qn("w:instr"), "NUMPAGES")
    paragraph._p.append(pages_field)


def configure_styles(document):
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(9.2)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.08

    heading = document.styles["Heading 1"]
    heading.font.name = "Arial"
    heading._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    heading._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    heading.font.size = Pt(10.5)
    heading.font.bold = True
    heading.font.color.rgb = TEAL
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(3)
    heading.paragraph_format.keep_with_next = True

    bullet = document.styles["List Bullet"]
    bullet.font.name = "Arial"
    bullet._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    bullet._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    bullet.font.size = Pt(8.9)
    bullet.font.color.rgb = INK
    bullet.paragraph_format.left_indent = Inches(0.2)
    bullet.paragraph_format.first_line_indent = Inches(-0.13)
    bullet.paragraph_format.space_before = Pt(0)
    bullet.paragraph_format.space_after = Pt(1.5)
    bullet.paragraph_format.line_spacing = 1.04


def build_cv():
    document = Document()
    document.core_properties.title = "Nicole Nikareayi - Graduate Business Analyst CV"
    document.core_properties.author = "Nicole Nikareayi"
    document.core_properties.subject = "Public job application CV"
    document.core_properties.keywords = (
        "Business Analyst, Graduate, Digital Product, Technology Innovation"
    )

    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)
    add_footer(section)
    configure_styles(document)

    name = document.add_paragraph()
    set_cell_free_paragraph(name, after=1)
    name_run = name.add_run("Nicole Nikareayi")
    set_run_font(name_run, 24, NAVY, bold=True)

    role = document.add_paragraph()
    set_cell_free_paragraph(role, after=4)
    role_run = role.add_run(
        "GRADUATE BUSINESS ANALYST | BUSINESS SYSTEMS, DATA & IMPLEMENTATION"
    )
    set_run_font(role_run, 11, TEAL, bold=True)

    contact = document.add_paragraph()
    set_cell_free_paragraph(contact, after=1)
    location = contact.add_run("Auckland, New Zealand  |  ")
    set_run_font(location, 9.2, MUTED)
    add_hyperlink(contact, "nigaray703@gmail.com", "mailto:nigaray703@gmail.com")
    separator = contact.add_run("  |  ")
    set_run_font(separator, 9.2, MUTED)
    add_hyperlink(
        contact,
        "linkedin.com/in/nikareayi-aisikaer",
        "https://www.linkedin.com/in/nikareayi-aisikaer",
    )
    separator = contact.add_run("  |  ")
    set_run_font(separator, 9.2, MUTED)
    add_hyperlink(
        contact,
        "github.com/nigaray703-ops",
        "https://github.com/nigaray703-ops",
    )

    rights = document.add_paragraph()
    set_cell_free_paragraph(rights, after=5)
    rights_run = rights.add_run(
        "Open Post-Study Work Visa | Full New Zealand work rights until 25 March 2029"
    )
    set_run_font(rights_run, 9.2, MUTED)

    add_section_heading(document, "Professional Summary")
    add_body(
        document,
        "Graduate business systems and technology professional with a Master of Technology Innovation in Business (Distinction) and practical experience across an approximately 400-hour academic internship, requirements and process analysis, application testing, data organisation, stakeholder documentation and privacy-conscious digital products.",
        after=4,
    )
    add_body(
        document,
        "Combines business, systems, data, AI and user-centred design capability. Targeting graduate and junior opportunities across business analysis, business systems, implementation, application support, data quality, reporting and continuous improvement.",
        after=4,
    )

    add_section_heading(document, "Core Capabilities")
    add_body(
        document,
        "Business analysis: requirements, user stories, process and workflow mapping, gap and option analysis. Systems and implementation: solution evaluation, testing, issue tracking, user guidance and documentation. Data and reporting: collection, cleaning, validation, reconciliation, Excel analysis, dashboards and structured records. Improvement: user research, accessibility testing, feedback synthesis and workflow refinement.",
        after=4,
    )

    add_section_heading(document, "Professional Experience")
    add_role_header(
        document,
        "Technology Innovation Intern - Business Analysis and Digital Delivery",
        "University of Waikato | Academic, A+",
        "Nov 2025 - Feb 2026 | Approximately 400 hours",
    )
    add_bullet(
        document,
        "Helped a student team deliver Ko ahau te awa, an interactive cultural storytelling prototype combining a map, puzzle interaction, responsive layouts and optional audio narration.",
    )
    add_bullet(
        document,
        "Translated project goals and feedback into user journeys, interaction flows, functional requirements and structured delivery tasks; owned the puzzle interaction from concept through refinement.",
    )
    add_bullet(
        document,
        "Compared Wix and Canva constraints, supported responsive Wix implementation across desktop and mobile, and applied clarity, accessibility and cognitive-load principles.",
    )
    add_bullet(
        document,
        "Supported testing, issue identification and repeated correction; maintained project records and helped deliver the final prototype, report and presentation on schedule.",
    )

    add_section_heading(document, "Selected Business Systems Projects")
    add_role_header(
        document,
        "Career Command Center - Bilingual Job Application Tracker",
        "Independent project",
        "Jun 2026 - Present",
    )
    add_bullet(
        document,
        "Defined business rules, data fields, status categories, workflows and information architecture for centralising applications, candidate portals and follow-up information.",
    )
    add_bullet(
        document,
        "Built dashboard metrics, conversion indicators, search, A-Z navigation, filtering, sorting, grouping and responsive record management.",
    )
    add_bullet(
        document,
        "Implemented Google authentication and user-scoped Supabase access for the private version; separated the public demo with fictional records, simulated sign-in and disabled cloud access.",
    )

    add_role_header(
        document,
        "Ana Tilim - Mobile-First Uyghur Language Learning Platform",
        "Independent project",
        "Jul 2026 - Present",
    )
    add_bullet(
        document,
        "Structured multilingual curriculum and course data for Uyghur script, ULY transliteration, vocabulary, grammar, practice activities and reading content.",
    )
    add_bullet(
        document,
        "Implemented RTL support, human-recorded audio, listening and dictation, offline progress, backup and restore, plus UID-scoped Supabase synchronisation with Row Level Security.",
    )
    add_bullet(
        document,
        "Created a verification runner covering course-data integrity, transliteration, audio manifests, interactions, authentication, cloud sync and full-content rendering across 464 interface states.",
    )

    document.add_page_break()

    add_section_heading(document, "Selected Data and Application Projects")
    add_role_header(
        document,
        "Harry Potter Knowledge Assistant - Curated AI Dataset",
        "COMPX500 | A+",
        "Aug 2025 - Sep 2025",
    )
    add_bullet(
        document,
        "Collected and integrated data from Wikipedia, Kaggle and Fandom using Python, BeautifulSoup, Pandas, Google Colab and Jupyter Notebook.",
    )
    add_bullet(
        document,
        "Cleaned schemas, removed duplicates, managed aliases, linked subjects and expanded the validated dataset to 16,245 rows.",
    )
    add_bullet(
        document,
        "Packaged a constrained upload bundle with a manifest and validation checks; designed source-aware, off-topic and uncertainty handling and compared cloud/local model results.",
    )

    add_role_header(
        document,
        "Cash Safety Assistant - Personal Finance Planning Web App",
        "Independent project",
        "Jun 2026 - Present",
    )
    add_bullet(
        document,
        "Translated cash-planning needs into editable profile and transaction models, recurring costs, income, safety thresholds and projected balances.",
    )
    add_bullet(
        document,
        "Implemented Firebase Authentication and UID-scoped Firestore storage, with a local Demo mode and explicit separation between local and authenticated cloud states.",
    )
    add_bullet(
        document,
        "Delivered a responsive Progressive Web App and a separate local SwiftUI companion prototype with JSON-backed on-device persistence.",
    )

    add_section_heading(document, "Leadership and Community")
    add_role_header(
        document,
        "Arts Committee Chair",
        "Dalian Jiaotong University",
        "Mar 2021 - Sep 2023",
    )
    add_bullet(
        document,
        "Planned campus events with 100+ participants and coordinated cross-functional teams across departments and student groups.",
    )
    add_role_header(
        document,
        "Ethnic Minority Liaison",
        "Dalian Jiaotong University",
        "Sep 2020 - Jul 2024",
    )
    add_bullet(
        document,
        "Supported communication and coordination for 30+ students and maintained structured records.",
    )

    add_section_heading(document, "Education")
    add_role_header(
        document,
        "Master of Technology Innovation in Business - Distinction",
        "University of Waikato, New Zealand",
        "2025 - 2026",
    )
    add_role_header(
        document,
        "Bachelor of Management (Accounting)",
        "Dalian Jiaotong University, China",
        "2020 - 2024",
    )

    add_section_heading(document, "Tools and Technical Skills")
    add_body(
        document,
        "Analysis and reporting: requirements analysis, process maps, functional specifications, test support, Excel cleaning, formulas, lookups, pivot tables, validation, reconciliation and dashboard metrics. Product and web: Figma, Miro, Wix, HTML, CSS, JavaScript, PWA, Electron, SwiftUI prototypes, Firebase, Firestore, Supabase and RLS. AI and data: Python, Pandas, Jupyter, Google Colab, BeautifulSoup, curated datasets and cloud/local LLM evaluation. Tools: Microsoft 365, GitHub, Jira and Confluence (basic), SQL and Power BI (foundational).",
        after=3,
    )

    add_section_heading(document, "Languages")
    add_body(
        document,
        "Uyghur - native | Mandarin Chinese - native | English - professional working proficiency | Korean - basic conversational | Te Reo Maori - introductory",
        after=2,
    )

    add_section_heading(document, "Additional Information")
    add_body(
        document,
        "Auckland-based | Full New Zealand driver licence | Available immediately and for business travel",
        after=0,
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_cv()
